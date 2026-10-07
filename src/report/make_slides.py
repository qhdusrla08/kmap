"""발표자료 PPTX 생성 (python-pptx). 수치는 outputs/tables, 그림은 outputs/figures에서 읽습니다.

블라인드 규정: 소속·로고를 넣지 않습니다. 팀명/성명은 --team 인자로만 넣습니다.
문서 속성(작성자·회사)은 비워 둡니다. PDF는 PowerPoint에서 '다른 이름으로 저장 → PDF'로 만듭니다.
사용: python -m src.report.make_slides --team "팀명"
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from src.config import ROOT

FIG = ROOT / "outputs" / "figures"
TAB = ROOT / "outputs" / "tables"
FONT = "맑은 고딕"
DARK = RGBColor(0x1F, 0x2A, 0x37)
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)


def j(name, default=None):
    p = TAB / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else (default or {})


def _style(run, size, bold=False, color=DARK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_slide(prs, title, bullets=(), image=None, table=None, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃
    W, H = prs.slide_width, prs.slide_height
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), W - Inches(1.0), Inches(0.8))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    _style(r, 26, True, ACCENT)
    left_w = W - Inches(1.0) if image is None and table is None else Inches(5.6)
    if bullets:
        bx = s.shapes.add_textbox(Inches(0.5), Inches(1.2), left_w, H - Inches(1.8))
        tf = bx.text_frame
        tf.word_wrap = True
        for k, b in enumerate(bullets):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            lvl = 1 if b.startswith("  ") else 0
            p.level = lvl
            rr = p.add_run()
            rr.text = ("· " if lvl else "■ ") + b.strip()
            _style(rr, 15 if lvl else 17)
            p.space_after = Pt(6)
    x0 = Inches(0.5) + (left_w if bullets else 0) + (Inches(0.2) if bullets else 0)
    avail_w = W - x0 - Inches(0.4)
    if image is not None and Path(image).exists():
        pic = s.shapes.add_picture(str(image), x0, Inches(1.2))
        scale = min(avail_w / pic.width, (H - Inches(1.6)) / pic.height)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    if table is not None:
        rows, cols = len(table), len(table[0])
        t = s.shapes.add_table(rows, cols, x0, Inches(1.2), avail_w, Emu(int(Inches(0.38)) * rows)).table
        for i, row in enumerate(table):
            for k, v in enumerate(row):
                c = t.cell(i, k)
                c.text = str(v)
                for p in c.text_frame.paragraphs:
                    for rr in p.runs:
                        _style(rr, 12, i == 0)
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


NAMES = {"hgb": "HGB 베이스라인", "hgb_v2": "HGB v2", "p2_coco": "YOLOv8s-P2 v1", "p2_scratch": "YOLOv8s-P2 scratch",
         "p3_coco": "YOLOv8s-P3 v1", "p2_coco_v2b": "YOLOv8s-P2 v2b", "p3_coco_v2b": "YOLOv8s-P3 v2b (최종)"}


def pct(v, d=1):
    return "—" if v is None or v != v else f"{v * 100:.{d}f}"


def compare_table(models):
    p = TAB / "model_compare_oof.csv"
    if not p.exists():
        return None
    import pandas as pd
    t = pd.read_csv(p)
    t = t[(t.subset == "all") & t.model.isin(models)]
    rows = [["모델", "AP50", "AP50:95", "R@FPPI0.1", "F1", "배경 과검 p99"]]
    for m in models:
        r = t[t.model == m]
        if len(r):
            r = r.iloc[0]
            rows.append([NAMES.get(m, m), f"{r.AP50:.3f}", f"{r.AP50_95:.3f}", f"{r.R_at_FPPI_0_1:.3f}", f"{r.F1_best:.3f}", f"{r.bg_max_p99:.3f}"])
    return rows


def leak_table(models):
    rows = [["모델", "T3 decoy 발화", "T4 삭제 후 발화", "T5 이식 recall", "T7 코어 차폐 recall", "α=0.1 원위치"]]
    import pandas as pd
    for m in models:
        L = j(f"leak_tests_{m}.json")
        if not L:
            continue
        st = TAB / f"stress_{m}_summary.csv"
        fl = None
        if st.exists():
            s_ = pd.read_csv(st)
            fl = float(s_[(s_.kind == "in_place") & (s_.alpha == 0.1)].recall.iloc[0])
        rows.append([NAMES.get(m, m), pct(L["T3"]["decoy_site_fire_rate"]) + "%", pct(L["T4"]["fire_rate_erased"]) + "%",
                     f'{L["T5"]["recall_transplant"]:.3f}', f'{L["T7"]["recall_core_occluded"]:.3f}', pct(fl) + "%"])
    return rows if len(rows) > 1 else None


def build(team, final="p3_coco_v2b"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cp = prs.core_properties
    cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.title = "X-ray 이물 탐지와 미탐지 조건 분석"
    t1 = j("leak_t1.json"); t2 = j("leak_t2.json")
    add_slide(prs, "X-ray 영상 기반 완제품 이물질 탐지 및 AI 미탐지 조건 분석",
              [f"제6회 K-인공지능 제조데이터 분석 경진대회 · 과제 ④", f"{team}"])
    add_slide(prs, "문제와 데이터의 함정", [
        "최종 X-ray 검사 단계: 미검 = 시장 유출 → Recall 최우선",
        "장비 NG 영상만 수집: 고유 2,532장, 라벨 500장(1,147객체), 음성 0장",
        "장비 색상 표시가 2,809장 전부에 있음, TXT 중심 99.8%가 표시 안",
        f"  표시만 쓰는 탐지기 F1 {t2.get('marking_only_dev_oof', {}).get('F1', 0.94):.2f} → 높은 점수 ≠ 이물 학습",
        "실제 이물: 2–4px 점, 띠 끝에 위치. 영상의 90%가 4시간 주기(감도점검 세션 추정)",
    ], FIG / "data_label_overlay.png")
    add_slide(prs, "표시 제거와 흔적-이물 독립화", [
        "팔레트 인덱스 ≥244 = 표시(손실 없이 정확히 검출)",
        "후광(바깥 약 3px) 실측 → 2px 팽창 + Telea + 밝기별 잡음",
        "decoy: 후광까지 합성한 가짜 표시(흔적 O, 이물 X)",
        "이식: Beer–Lambert 투과율(이물 O, 흔적 X)",
        "평가 입력은 clean 영상(decoy 없음)",
    ], FIG / "qc_inpaint_methods.png")
    add_slide(prs, "분할과 평가 설계", [
        "그룹 = 달력 날짜(전 호기 공통) — 같은 초 다호기 세션·시험편 반복 촬영 대응",
        "동결 테스트 100장(7개 날짜, 4개 월, 3개 호기) — 마지막에 한 번만 평가",
        "dev 400장 날짜 4-fold OOF로 선택·보정·임계값 결정",
        "음성 0장 → 같은 제품 배경의 실제 픽셀로 음성 대용(bg_max), FPPI·유병률 환산",
        "추가: 호기 hold-out(LOMO), 전향(6–7월 → 8–9월)",
    ], FIG / "data_time_structure.png")
    models = [m for m in ("hgb", "p2_scratch", "p2_coco", "p3_coco", "p2_coco_v2b", "p3_coco_v2b")]
    add_slide(prs, "모델 비교 (dev OOF, IoU 0.5)", [
        "IoU 0.5 정확도는 고전 모델도 포화 → 차이는 강건성·박스 정밀도·속도에서",
        "COCO 사전학습: scratch 대비 AP50 +1.4%p, 배경 과검 감소",
    ], table=compare_table(models))
    add_slide(prs, "지름길(표시 누수) 검증", [
        "T3 decoy 흔적 자리, T4 이물만 지운 자리에서 발화하지 않아야 함",
        "T5 흔적 없는 위치로 옮긴 이물도 찾아야 함",
        "T7 코어를 가리면 무너지고 링을 가려도 유지 → 이물 자체를 봄",
        "α=0.1 원위치 바닥값: v2(흔적O·이물X 음성)로 제거",
    ], table=leak_table(models))
    add_slide(prs, "조건 분석: 대비 × 위치 탐지 한계", [
        "실제 라벨은 '띠 끝의 고대비 점'뿐 → 통제 스트레스 시험",
        "투과율 감쇠 T_α = 1 − α(1 − T), 흔적 없는 띠 끝·내부·가장자리 이식",
    ], FIG / "report_stress.png")
    add_slide(prs, "보정과 3구간 판정", [
        "Platt 보정: ECE 0.216 → 0.002 (최종모델)",
        "통과 < t_low = min(사전 0.947, 검출 임계값 0.833) — 검출하면 자동 통과 없음",
        "OOD 가드: 제품 구조가 학습 범위 밖이면 재검사",
        "배출 ≥ t_high (음성 대용 400장 중 0장)",
        "TTA 불확실성(std ≥ 0.1) → 통과 구간이라도 재검사",
    ], FIG / "report_zones.png")
    H = j("../predictions/holdout_metrics.json")
    if H and final in H:
        r = H[final]
        add_slide(prs, "동결 테스트 결과 (한 번만 평가)", [
            f"AP50 {r['AP50']:.3f} · AP50:95 {r['AP50_95']:.3f} · R@FPPI0.1 {r['R_at_FPPI_0_1']:.3f}",
            f"OOF 임계값에서 P {r['at_oof_F1_threshold']['precision']:.3f} · R {r['at_oof_F1_threshold']['recall']:.3f} · F1 {r['at_oof_F1_threshold']['F1']:.3f}",
            f"객체 recall 95% CI(날짜 부트스트랩) {r['object_recall_at_oof_thr_CI95'][0]:.3f}–{r['object_recall_at_oof_thr_CI95'][1]:.3f}",
            f"사전 규칙 3구간: 1객체 61장 중 {r['one_obj_pass']}장 통과 → 이식 실패(다음 장)",
        ] + ([f"사후 확인(설계 전 결과를 본 데이터): 현행 규칙 1객체 통과 {r['posthoc_revised_rule']['one_obj_pass']}/{r['posthoc_revised_rule']['n_one_obj']}"]
             if r.get("posthoc_revised_rule") else []))
    TT = j(f"threshold_transfer_{final}.json")
    if TT:
        pooled = {(d["mode"], d["rule"]): d for d in TT["pooled"]}
        def cell(mode, rule):
            d = pooled.get((mode, rule))
            return f"{d['pass_one']}/{d['n_one']} · {100 * d['neg_flag_rate']:.2f}%" if d else "—"
        modes = [m for m in ("lofo", "nested") if any(k[0] == m for k in pooled)]
        rows = [["규칙"] + [f"{m.upper()} 누수 · 비용" for m in modes]]
        names = {"R0": "R0 사전(recall 0.99 최대 t)", "R1": "R1 min(R0, 검출 임계값) ← 선택", "R2": "R2 호기별 R0",
                 "R3": "R3 부트스트랩 5%", "R4": "R4 비용 ≤0.5% 최소 t"}
        for r in ("R0", "R1", "R2", "R3", "R4"):
            rows.append([names[r]] + [cell(m, r) for m in modes])
        add_slide(prs, "통과 임계값 이식 실패: 원인과 dev 검증", [
            "테스트 통과 9장 = 모두 3호기 1띠 영상, 탐지는 됨(점수 0.45–0.54)",
            "원인: R0는 양성 분포 아래 끝에 붙어 여유 0 + 3호기 1띠 점수 척도 하락 + 재학습 모델 이동",
            "dev 시뮬레이션(LOFO·중첩 학습)으로 재현, 기준은 실행 전 고정(비용 ≤ 1%)",
        ], table=rows)
        add_slide(prs, "t_low 위치별 미검 누수 vs 재검사 비용", [
            "음성 대용 비용은 원점수 0.1–0.5에서 거의 평평, 양성은 0.54 이상",
            "양성 쪽에서 정한 규칙(R0–R3)은 여유가 없고, 비용 쪽 규칙(R4)은 누수 0 대신 비용 증가",
        ], FIG / f"threshold_tradeoff_{final}.png")
    add_slide(prs, "모니터링 정정: 3호기 '드리프트'의 실체", [
        "영상당 top-1 박스로 집계 → 이물 대비·영상 품질은 전 기간 안정(장비 열화 신호 없음)",
        "3호기가 1띠 제품으로 바뀐 07-21부터 점수만 0.69 → 0.58–0.67 (1·2호기 1띠는 0.72–0.73)",
        "통과 판정 3호기 44장 전부 검출 임계값 이상 박스 1개 → 탐지 실패가 아닌 임계값 문제",
        "1호기 07-27 띠 없는 다른 제품(OOD) 93장 → OOD 가드로 재검사",
    ], FIG / "monitoring_trend.png")
    sp = j("speed.json")
    add_slide(prs, "현장 활용", [
        "장비 NG 영상의 2차 판정기: 통과 / 재검사 / 배출",
        f"처리 시간: 전처리 {sp.get('preprocess_ms', {}).get('mean', float('nan')):.1f}ms + 탐지 {sp.get('yolo_cuda_ms', {}).get('mean', float('nan')):.1f}ms/장 (GPU)",
        "점검 세션별 이물 대비(장비)·top-1 점수(판정 척도)를 분리 감시 → 열화·재보정 알림",
        "도입 전 OK 영상 수집으로 미검 감소 효과 검증 필요",
    ])
    add_slide(prs, "재현성과 규정 준수", [
        "python run_all.py 단일 명령 (core / full 프로필), 고정 seed·동결 분할 해시",
        "early stopping 없음, last.pt OOF, 테스트는 final 단계에서 한 번",
        "외부 자원: COCO 사전학습 YOLOv8s(AGPL-3.0) — scratch 비교로 기여 정량화",
        "블라인드 점검 스크립트로 제출물의 식별 정보·절대경로 검사",
    ])
    add_slide(prs, "한계와 다음 단계", [
        "양품(OK) 영상 0장: 이미지 정밀도·현장 과검률은 대용 추정",
        "시험편 반복 촬영 가능성: 새로운 이물에 대한 일반화는 미보장",
        "공식 정답 TXT의 누락 사례 존재(마킹 3개·TXT 2개)",
        "현장 적용 시 표시 없는 원영상 수집이 가장 큰 개선 요인",
        "3호기 잡음 정합 증강(ablation): hold-out F1 0.761→0.825, 분포 안 3호기 R@FPPI 0.958→0.518 → 미채택",
        "통과 규칙 변경은 테스트 실패 후 설계·dev 검증 → 현장 전향 검증 필요",
    ])
    return prs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--team", default="팀명")
    ap.add_argument("--out", default="reports/slides.pptx")
    ap.add_argument("--final", default="p3_coco_v2b")
    a = ap.parse_args()
    prs = build(a.team, a.final)
    out = ROOT / a.out
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
